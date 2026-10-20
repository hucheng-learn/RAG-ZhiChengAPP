import asyncio
import base64
import json
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from docx import Document

from common.config import Settings
from tests.api_workflow_tests.test_workflows import course, wait_task
from utils.parsers.question_recognition import recognize_questions

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jU1sAAAAASUVORK5CYII="
)


def send(client, files, **fields):
    return client.post(
        "/api/v1/documents/import-folder",
        data={"content_type": "lecture", "course_id": "python", **fields},
        files=[("files", (name, body)) for name, body in files],
    )


def test_default_limits_and_missing_mineru(system):
    client, _ = system
    assert Settings(_env_file=None).max_upload_bytes == 200 * 1024 * 1024
    limits = client.get("/api/v1/capabilities").json()["upload_limits"]
    assert limits["max_file_bytes"] == 200 * 1024 * 1024
    assert ".pdf" in limits["supported_extensions"] and not limits["mineru"]
    response = client.post(
        "/api/v1/documents/import",
        data={"content_type": "lecture"},
        files={"file": ("a.pdf", b"pdf")},
    )
    assert response.status_code == 503 and response.json()["code"] == "mineru_unavailable"


@pytest.mark.parametrize("reference", [
    "![figure](../images/a.png)",
    '<img src="../images/a.png" style="zoom: 67%;" />',
])
def test_folder_associates_reference_images_preserves_paths_and_deletes_assets(system, reference):
    client, runtime = system
    course(client)
    response = send(
        client,
        [
            ("root/lessons/a.md", f"# Lesson\n{reference}".encode()),
            ("root/images/a.png", PNG),
            ("root/notes.bin", b"skip"),
        ],
    )
    assert response.status_code == 202, response.text
    payload = response.json()
    assert payload["attachments"] == ["root/images/a.png"]
    assert len(payload["items"]) == 2 and payload["items"][1]["skipped"]
    task = payload["items"][0]["task"]
    assert wait_task(client, task["task_id"])["status"] == "completed"
    version = asyncio.run(runtime.repo.get("versions", task["version_id"]))
    assert version["source_path"] == "root/lessons/a.md"
    url = f"/api/v1/documents/{task['document_id']}/asset"
    image = client.get(url, params={"path": "../images/a.png", "version_id": task["version_id"]})
    assert image.status_code == 200 and image.content == PNG
    assert (
        client.get(
            url, params={"path": "../../private.png", "version_id": task["version_id"]}
        ).status_code
        == 404
    )
    deleted = client.delete(f"/api/v1/documents/{task['document_id']}").json()
    assert wait_task(client, deleted["task_id"])["status"] == "completed"
    assert not runtime.blobs.objects


def test_folder_same_basename_is_distinct_and_attachment_changes_break_reuse(system):
    client, _ = system
    course(client)
    first = send(client, [("root/a/file.md", b"same"), ("root/b/file.md", b"same")]).json()
    tasks = [item["task"] for item in first["items"]]
    assert len({task["document_id"] for task in tasks}) == 2
    for task in tasks:
        assert wait_task(client, task["task_id"])["status"] == "completed"
    source = [("root/a.md", b"![a](a.png)"), ("root/a.png", PNG)]
    original = send(client, source).json()["items"][0]["task"]
    assert wait_task(client, original["task_id"])["status"] == "completed"
    changed = send(client, [source[0], ("root/a.png", PNG + b"change")]).json()["items"][0]["task"]
    assert changed["document_id"] != original["document_id"]


@pytest.mark.parametrize("name", ["../a.md", "/a.md", "C:/a.md", "root/../a.md"])
def test_folder_rejects_unsafe_paths(system, name):
    client, _ = system
    assert send(client, [(name, b"hello")]).status_code == 400


def test_word_question_bank_uses_existing_fields_without_markdown_only_gate(system):
    client, _ = system
    course(client)
    doc = Document()
    doc.add_heading("Python", level=2)
    doc.add_paragraph("题库编码: python")
    doc.add_heading("q001", level=3)
    doc.add_paragraph("题干: 列表是否可变？")
    doc.add_paragraph("答案: 是")
    data = BytesIO()
    doc.save(data)
    response = send(client, [("questions.docx", data.getvalue())], content_type="question_bank")
    task = response.json()["items"][0]["task"]
    assert wait_task(client, task["task_id"])["status"] == "completed"
    questions = client.get("/api/v1/questions").json()["items"]
    assert len(questions) == 1 and "answer" not in questions[0]
    assert questions[0]["source"]["format"] == "normalized"


def test_word_embedded_images_are_preserved_without_mineru(system):
    client, runtime = system
    course(client)
    doc = Document()
    doc.add_paragraph("图片资料")
    doc.add_picture(BytesIO(PNG))
    data = BytesIO()
    doc.save(data)
    task = send(client, [("a.docx", data.getvalue())]).json()["items"][0]["task"]
    assert wait_task(client, task["task_id"])["status"] == "completed"
    version = asyncio.run(runtime.repo.get("versions", task["version_id"]))
    assert len(version["assets"]) == 1
    assert "word/media/image1.png" in version["assets"][0]["path"]


def test_mineru_pdf_enters_existing_pipeline_and_downloads_as_pdf(system):
    client, runtime = system
    course(client)
    runtime.mineru.require = lambda: None
    runtime.mineru.parse = AsyncMock(return_value=("# PDF内容\n列表是可变集合", {}))
    task = send(client, [("a.pdf", b"%PDF-test")]).json()["items"][0]["task"]
    assert wait_task(client, task["task_id"])["status"] == "completed"
    response = client.get(f"/api/v1/documents/{task['document_id']}/source")
    assert response.headers["content-type"] == "application/pdf"


def test_unstructured_questions_keep_original_answers_and_reject_invention():
    source = "1. 列表是否可变？\n答案：是\n解析：支持修改。"
    row = {
        "start_line": 1,
        "end_line": 3,
        "stem": "列表是否可变？",
        "answer": "是",
        "explanation": "支持修改。",
    }

    class Model:
        client = object()

        async def complete(self, messages):
            return json.dumps({"questions": [row], "incomplete": False}, ensure_ascii=False)

    questions = asyncio.run(recognize_questions(source, "d", "v", "c", Model()))
    assert questions[0]["answer"] == "是" and questions[0]["review_status"] == "unreviewed"
    row["answer"] = "不存在的答案"
    with pytest.raises(ValueError, match="原文"):
        asyncio.run(recognize_questions(source, "d", "v", "c", Model()))


def test_batch_limits_and_per_file_failures_do_not_drop_valid_files(system):
    client, runtime = system
    course(client)
    runtime.settings.max_upload_bytes = 1024
    response = send(client, [("large.txt", b"x" * 1025)])
    assert response.status_code == 413
    runtime.settings.max_batch_bytes = 1024
    assert send(client, [("a.txt", b"x" * 600), ("b.txt", b"x" * 600)]).status_code == 413
    response = send(client, [("a.md", b"valid"), ("bad.md", b"\xff"), ("a.pdf", b"pdf")])
    items = response.json()["items"]
    assert items[0]["task"] and items[1]["code"] == "invalid_encoding"
    assert items[2]["code"] == "mineru_unavailable"


def test_durable_queue_accepts_more_than_in_memory_slots(system):
    client, runtime = system
    course(client)
    client.portal.call(runtime.imports.close)
    response = send(client, [(f"{i}.md", f"document {i}".encode()) for i in range(12)])
    tasks = [item["task"] for item in response.json()["items"]]
    assert len(tasks) == 12 and runtime.imports.queue.qsize() == 8
    client.portal.call(runtime.imports.start)
    for task in tasks:
        assert wait_task(client, task["task_id"])["status"] == "completed"


@pytest.mark.parametrize("reference", ["![图](a.png)", '<img src="a.png" />'])
def test_mineru_markdown_image_ocr_remains_in_original_section(system, reference):
    from pydantic import SecretStr

    client, runtime = system
    course(client)
    runtime.settings.mineru_api_token = SecretStr("mock")
    runtime.mineru.parse = AsyncMock(return_value=("图片里有文字", {}))
    task = send(
        client,
        [("a.md", f"# 第一节\n{reference}\n\n# 第二节\n普通文字".encode()), ("a.png", PNG)],
    ).json()["items"][0]["task"]
    assert wait_task(client, task["task_id"])["status"] == "completed"
    sections = asyncio.run(runtime.repo.find("sections", {"document_id": task["document_id"]}))
    first = next(section for section in sections if section["title"] == "第一节")
    second = next(section for section in sections if section["title"] == "第二节")
    assert "图片里有文字" in first["body"] and "图片里有文字" not in second["body"]
    assert first["source"]["format"] == "normalized"


def test_question_images_are_versioned_and_solution_images_stay_out_of_practice(system):
    client, _ = system
    course(client)
    source = """## 题库
- 题库编码: pictures
### q1
- 题干: 看图回答 ![题图][stem]
- 答案: A ![答案图](answer.png)

[stem]: stem.png
"""
    task = send(
        client,
        [("bank.md", source.encode()), ("stem.png", PNG), ("answer.png", PNG)],
        content_type="question_bank",
    ).json()["items"][0]["task"]
    assert wait_task(client, task["task_id"])["status"] == "completed"
    question = client.get("/api/v1/questions").json()["items"][0]
    assert question["document_id"] == task["document_id"]
    assert question["version_id"] == task["version_id"]
    assert question["images"] == ["stem.png"]
    assert "solution_images" not in question and "answer" not in question
    answer = client.get(f"/api/v1/questions/{question['question_id']}/solution").json()
    assert answer["images"] == ["answer.png"]


def test_unstructured_word_falls_back_to_model_in_full_import_pipeline(system):
    client, runtime = system
    course(client)
    doc = Document()
    doc.add_paragraph("1. 列表是否可变？")
    doc.add_paragraph("答案：是")
    data = BytesIO()
    doc.save(data)
    runtime.llm.complete = AsyncMock(
        return_value=json.dumps(
            {
                "questions": [
                    {"start_line": 1, "end_line": 3, "stem": "列表是否可变？", "answer": "是"}
                ],
                "incomplete": False,
            },
            ensure_ascii=False,
        )
    )
    task = send(client, [("bank.docx", data.getvalue())], content_type="question_bank").json()[
        "items"
    ][0]["task"]
    assert wait_task(client, task["task_id"])["status"] == "completed"
    question = client.get("/api/v1/questions").json()["items"][0]
    assert question["stem"] == "列表是否可变？" and "answer" not in question
    assert runtime.llm.complete.await_count == 1

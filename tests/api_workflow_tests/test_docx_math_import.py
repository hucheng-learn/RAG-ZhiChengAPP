"""检查本地 Word 公式经过正式导入后仍进入标准化原文、切片和索引。"""

import asyncio

from docx import Document

import services.imports as imports
from tests.api_workflow_tests.test_workflows import course, wait_task
from tests.component_tests.test_docx_math import append_math, word_bytes


def formula_document():
    document = Document()
    document.add_heading("分数", 1)
    paragraph = document.add_paragraph("计算")
    append_math(
        paragraph,
        "<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num>"
        "<m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>",
    )
    return word_bytes(document)


def submit(client, data, document_id=None):
    path = (
        f"/api/v1/documents/{document_id}/versions" if document_id else "/api/v1/documents/import"
    )
    response = client.post(
        path,
        data={"content_type": "lecture", "course_id": "python"},
        files={"file": ("formula.docx", data)},
    )
    assert response.status_code == 202, response.text
    return response.json()


def test_local_docx_formula_survives_normalization_chunks_and_embedding(system):
    client, runtime = system
    course(client)
    data = formula_document()
    task = submit(client, data)
    assert wait_task(client, task["task_id"])["status"] == "completed"
    version = asyncio.run(runtime.repo.get("versions", task["version_id"]))
    assert version["document_parser_version"] == "docx-v3-omml"
    normalized = asyncio.run(runtime.blobs.get(version["normalized_key"])).decode()
    assert "$\\frac{1}{2}$" in normalized
    chunks = asyncio.run(runtime.repo.find("chunks", {"version_id": task["version_id"]}))
    assert any("\\frac{1}{2}" in chunk["text"] for chunk in chunks)
    assert any("\\frac{1}{2}" in chunk["embedding_text"] for chunk in chunks)
    assert any(row["version_id"] == task["version_id"] for row in runtime.vectors.rows.values())
    assert submit(client, data)["reused"] is True


def test_parser_upgrade_does_not_reuse_old_formula_import(system, monkeypatch):
    client, runtime = system
    course(client)
    data = formula_document()
    current = imports.DOCX_PARSER_VERSION
    monkeypatch.setattr(imports, "DOCX_PARSER_VERSION", "docx-v2")
    first = submit(client, data)
    assert wait_task(client, first["task_id"])["status"] == "completed"
    monkeypatch.setattr(imports, "DOCX_PARSER_VERSION", current)
    updated = submit(client, data, first["document_id"])
    assert not updated["reused"] and updated["version_id"] != first["version_id"]
    assert wait_task(client, updated["task_id"])["status"] == "completed"
    active = asyncio.run(runtime.repo.get("documents", first["document_id"]))
    assert active["active_version"] == updated["version_id"]


def test_docx_parser_change_preserves_markdown_fingerprint(system, monkeypatch):
    _, runtime = system
    markdown = runtime.imports.fingerprint("lecture", source_name="lecture.md")
    word = runtime.imports.fingerprint("lecture", source_name="lecture.docx")
    monkeypatch.setattr(imports, "DOCX_PARSER_VERSION", "docx-future")
    assert runtime.imports.fingerprint("lecture", source_name="lecture.md") == markdown
    assert runtime.imports.fingerprint("lecture", source_name="lecture.docx") != word


def test_unsupported_formula_update_keeps_previous_active_version(system):
    client, runtime = system
    course(client)
    original = submit(client, formula_document())
    assert wait_task(client, original["task_id"])["status"] == "completed"
    document = Document()
    append_math(document.add_paragraph(), "<m:unknown><m:r><m:t>x</m:t></m:r></m:unknown>")
    updated = submit(client, word_bytes(document), original["document_id"])
    failure = wait_task(client, updated["task_id"])
    assert failure["status"] == "failed" and "公式" in failure["error"]
    active = asyncio.run(runtime.repo.get("documents", original["document_id"]))
    assert active["active_version"] == original["version_id"]

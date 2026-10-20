"""检查 Markdown/Word 标题、代码块、表格和正文的解析边界。"""

from io import BytesIO

from docx import Document

from utils.parsers.docx import parse_docx
from utils.parsers.markdown import parse_markdown


def md(source):
    return parse_markdown(source, document_id="d", version_id="v")


def test_setext_code_nested_headings_and_duplicate_names():
    source = "前言\n\nPython\n======\n正文\n\n```py\n# 不是标题\n```\n\n## 变量\nx\n\n## 变量\ny\n"
    sections = md(source).sections
    assert [s.title for s in sections] == ["", "Python", "变量", "变量"]
    assert sections[2].parent_section_id == sections[1].section_id
    assert sections[2].section_id != sections[3].section_id
    assert sections[3].title_path == ["Python", "变量"]
    assert "# 不是标题" in sections[1].body
    assert sections[1].source.heading_start == 3
    lines = source.splitlines(keepends=True)
    for section in sections:
        location = section.source
        assert section.body == "".join(lines[location.start - 1 : location.end_exclusive - 1])


def test_quoted_and_list_headings_are_body_and_table_is_preserved():
    result = md("# 标题\n> ## 引用\n\n- ## 列表\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    assert len(result.sections) == 1
    assert "| 1 | 2 |" in result.sections[0].body
    assert "> ## 引用" in result.sections[0].body


def test_empty_no_heading_and_heading_at_eof():
    assert md("").sections[0].body == ""
    assert md("正文").sections[0].body == "正文"
    assert md("# 标题").sections[0].source.start == 2
    assert md("# 标题").sections[0].body == ""


def test_docx_block_order_and_hierarchy():
    document = Document()
    document.add_heading("Python", level=1)
    document.add_paragraph("表格前")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "值"
    document.add_paragraph("表格后")
    document.add_heading("变量", level=2)
    document.add_paragraph("赋值")
    data = BytesIO()
    document.save(data)
    result = parse_docx(data.getvalue(), document_id="d", version_id="v")
    assert result.sections[0].body == "表格前\n\n名称 | 值\n\n表格后"
    assert result.sections[0].source.start == 2
    assert result.sections[0].source.end_exclusive == 5
    assert result.sections[1].parent_section_id == result.sections[0].section_id
    assert result.sections[1].title_path == ["Python", "变量"]

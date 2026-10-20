"""从真实 DOCX XML 回归行内、独立和表格公式，不以扁平文本冒充公式。"""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from common.errors import BusinessError
from utils.parsers.docx import parse_docx


def append_math(paragraph, expression, *, display=False):
    root = f"<m:oMath>{expression}</m:oMath>"
    if display:
        root = f"<m:oMathPara>{root}</m:oMathPara>"
    paragraph._p.append(parse_xml(f"<w:p {nsdecls('w', 'm')}>{root}</w:p>")[0])


def word_bytes(document):
    data = BytesIO()
    document.save(data)
    return data.getvalue()


def parsed(document):
    return parse_docx(word_bytes(document), document_id="d", version_id="v")


def test_docx_inline_math_preserves_surrounding_text_and_symbols():
    document = Document()
    document.add_heading("梯度下降", 1)
    paragraph = document.add_paragraph("这里的")
    append_math(paragraph, "<m:r><m:t>∇L</m:t></m:r>")
    paragraph.add_run("是梯度；")
    append_math(paragraph, "<m:r><m:t>α</m:t></m:r>")
    paragraph.add_run("是学习率。")
    body = parsed(document).sections[0].body
    assert "\\nabla" in body and "\\alpha" in body
    assert body.startswith("这里的$") and body.endswith("$是学习率。")
    assert body.index("\\nabla") < body.index("是梯度") < body.index("\\alpha")


def test_docx_display_formula_is_not_empty_and_has_original_block_location():
    document = Document()
    document.add_heading("更新公式", 1)
    paragraph = document.add_paragraph()
    append_math(
        paragraph,
        "<m:sSub><m:e><m:r><m:t>θ</m:t></m:r></m:e>"
        "<m:sub><m:r><m:t>k+1</m:t></m:r></m:sub></m:sSub>"
        "<m:r><m:t>=θ−α∇L</m:t></m:r>",
        display=True,
    )
    document.add_paragraph("下一段")
    assert paragraph.text == ""  # 原缺陷：python-docx 的 text 不包含 OMML。
    section = parsed(document).sections[0]
    assert "$$\n" in section.body and "_{k+1}" in section.body
    assert "\\theta" in section.body and "\\alpha" in section.body
    assert section.body.index("$$") < section.body.index("下一段")
    assert section.source.start == 2 and section.source.end_exclusive == 4


@pytest.mark.parametrize(
    ("expression", "expected"),
    [
        (
            "<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>",
            "\\frac{a}{b}",
        ),
        (
            "<m:rad><m:deg><m:r><m:t>3</m:t></m:r></m:deg>"
            "<m:e><m:r><m:t>x</m:t></m:r></m:e></m:rad>",
            "\\sqrt[3]{x}",
        ),
        (
            "<m:sSubSup><m:e><m:r><m:t>x</m:t></m:r></m:e>"
            "<m:sub><m:r><m:t>i</m:t></m:r></m:sub>"
            "<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSubSup>",
            "{x}_{i}^{2}",
        ),
        (
            "<m:m><m:mr><m:e><m:r><m:t>1</m:t></m:r></m:e>"
            "<m:e><m:r><m:t>2</m:t></m:r></m:e></m:mr></m:m>",
            "\\begin{matrix}\n1 & 2",
        ),
    ],
)
def test_docx_structured_math(expression, expected):
    document = Document()
    append_math(document.add_paragraph(), expression)
    assert expected in parsed(document).sections[0].body


def test_docx_heading_and_table_formulas_preserve_hierarchy_and_order():
    document = Document()
    heading = document.add_heading("参数", 1)
    append_math(heading, "<m:r><m:t>α</m:t></m:r>")
    document.add_paragraph("表格前")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "比率"
    append_math(
        table.cell(0, 1).paragraphs[0],
        "<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num>"
        "<m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>",
    )
    document.add_paragraph("表格后")
    result = parsed(document)
    section = result.sections[0]
    assert "\\alpha" in section.title and section.title_path == [section.title]
    assert section.body == "表格前\n\n比率 | $\\frac{1}{2}$\n\n表格后"
    assert section.source.start == 2 and section.source.end_exclusive == 5


def test_docx_math_keeps_tabs_breaks_and_normal_hyperlink_results():
    document = Document()
    paragraph = document.add_paragraph("前\t")
    append_math(paragraph, "<m:r><m:t>x</m:t></m:r>")
    paragraph.add_run("\n后")
    paragraph._p.append(
        parse_xml(f'<w:hyperlink {nsdecls("w")}><w:r><w:t>链接文字</w:t></w:r></w:hyperlink>')
    )
    assert parsed(document).sections[0].body == "前\t$x$\n后链接文字"


def test_unsupported_nested_formula_fails_instead_of_silently_losing_content():
    document = Document()
    append_math(
        document.add_paragraph(),
        "<m:r><m:t>x+</m:t></m:r><m:unknown><m:r><m:t>y</m:t></m:r></m:unknown>",
    )
    with pytest.raises(BusinessError) as error:
        parsed(document)
    assert error.value.code == "docx_formula_unsupported"


def test_pagination_markers_inside_math_do_not_hide_equation_text():
    document = Document()
    append_math(
        document.add_paragraph(),
        "<m:r><m:t>x</m:t><w:lastRenderedPageBreak/><w:br/><m:t>+y</m:t></m:r>",
    )
    assert parsed(document).sections[0].body == "$x+y$"

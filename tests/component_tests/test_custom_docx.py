"""检查自定义 Word 标题样式能否保留章节层级。"""

from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from utils.parsers.docx import parse_docx


def test_explicit_custom_style_mapping_preserves_hierarchy():
    document = Document()
    document.styles.add_style("一级标题", WD_STYLE_TYPE.PARAGRAPH)
    document.styles.add_style("二级标题", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Python", "一级标题")
    document.add_paragraph("列表", "二级标题")
    document.add_paragraph("可变有序集合")
    buffer = BytesIO()
    document.save(buffer)
    result = parse_docx(
        buffer.getvalue(),
        document_id="d",
        version_id="v",
        heading_styles={"一级标题": 1, "二级标题": 2},
    )
    assert result.sections[1].title_path == ["Python", "列表"]
    assert result.sections[1].parent_section_id == result.sections[0].section_id
    assert any("显式配置" in warning for warning in result.warnings)

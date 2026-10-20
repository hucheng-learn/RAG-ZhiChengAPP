import re
from io import BytesIO

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from schemas.document import ParsedDocument, Section, SourceLocation
from utils.parsers.docx_math import paragraph_text

DOCX_PARSER_VERSION = "docx-v3-omml"


def parse_docx(
    source: bytes,
    *,
    document_id: str,
    version_id: str,
    heading_styles: dict[str, int] | None = None,
    include_images: bool = False,
) -> ParsedDocument:
    """按原件顶层块顺序解析；标题用样式识别，源位置使用段落/表格编号而非页码。"""
    document = Document(BytesIO(source))
    blocks = list(document.iter_inner_content())
    warnings = []
    sections: list[Section] = []
    stack: list[Section] = []
    body: list[str] = []
    title, level, body_start, heading_start = "", 0, 1, None
    custom_style_used = False

    def flush(end: int):
        nonlocal body
        while stack and stack[-1].heading_level >= level:
            stack.pop()
        parent = stack[-1] if stack else None
        section = Section(
            section_id=f"{document_id}:{version_id}:s{len(sections)}",
            parent_section_id=parent.section_id if parent else None,
            document_id=document_id,
            version_id=version_id,
            section_index=len(sections),
            title=title,
            title_path=([*parent.title_path, title] if parent else [title]) if title else [],
            heading_level=level,
            body="\n\n".join(body),
            source=SourceLocation(
                format="docx",
                unit="block",
                start=body_start,
                end_exclusive=end,
                heading_start=heading_start,
            ),
        )
        sections.append(section)
        if level:
            stack.append(section)
        body = []

    for number, block in enumerate(blocks, start=1):
        if isinstance(block, Paragraph):
            style_id = block.style.style_id if block.style else ""
            match = re.fullmatch(r"Heading([1-9])", style_id)
            heading_level = int(match[1]) if match else None
            # 优先使用 DOCX 自带的大纲级别；可读样式名称映射仅作明确配置的兜底。
            outline = block._p.xpath("./w:pPr/w:outlineLvl/@w:val")
            style = block.style
            seen = set()
            while not outline and style is not None and style.style_id not in seen:
                seen.add(style.style_id)
                outline = style.element.xpath("./w:pPr/w:outlineLvl/@w:val")
                style = style.base_style
            if outline and 0 <= int(outline[0]) <= 8:
                heading_level = int(outline[0]) + 1
            if heading_level is None and block.style and heading_styles:
                heading_level = heading_styles.get(block.style.name)
                custom_style_used |= heading_level is not None
            if heading_level is not None:
                if body or heading_start is not None:
                    flush(number)
                title, level = paragraph_text(block), heading_level
                body_start, heading_start = number + 1, number
            else:
                body.append(paragraph_text(block))
        elif isinstance(block, Table):
            # 保留原顺序，表格转换为标准化文本；块编号指向原 DOCX 表格。
            body.append(
                "\n".join(
                    " | ".join(
                        "\n".join(paragraph_text(p) for p in cell.paragraphs)
                        for cell in row.cells
                    )
                    for row in block.rows
                )
            )
        if include_images:
            element = block._p if isinstance(block, Paragraph) else block._tbl
            for rid in element.xpath(".//a:blip/@r:embed"):
                part = document.part.related_parts.get(rid)
                if part is not None:
                    body.append(f"![文档图片]({str(part.partname).lstrip('/')})")
    flush(len(blocks) + 1)
    if not any(section.heading_level for section in sections):
        warnings.append("未识别到标准标题样式，请核对章节结构")
    if custom_style_used:
        warnings.append("已按显式配置映射自定义标题样式，请核对层级")
    if not include_images and document.part.element.xpath(".//w:drawing | .//w:pict"):
        warnings.append("文档包含图片，本阶段未识别图片内容")
    return ParsedDocument(parser_version=DOCX_PARSER_VERSION, sections=sections, warnings=warnings)


def docx_images(source):
    """只读取嵌入的图片 part，不访问外部关系。"""
    document = Document(BytesIO(source))
    return {
        str(part.partname).lstrip("/"): part.blob
        for part in document.part.related_parts.values()
        if part.content_type.startswith("image/")
    }
